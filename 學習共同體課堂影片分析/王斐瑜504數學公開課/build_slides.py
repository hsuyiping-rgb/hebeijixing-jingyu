"""由 output/簡報/slides_content.json 產生 learning 風格 16:9 簡報（PPTX＋HTML）。

各段的完整描述—詮釋—理答—反思（analysis.txt）放進 PPTX 的備忘稿，投影片上只放濃縮版。
用法：python -X utf8 -I build_slides.py <lesson_dir> [輸出檔名.pptx]
"""
import html, json, re, sys
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

root = Path(sys.argv[1])
out = root / 'output'
deck_dir = out / '簡報'
C = json.loads((deck_dir / 'slides_content.json').read_text(encoding='utf-8'))
board_img = out / 'screenshots' / 'board_31m45.png'

BG, GREEN, BODY, GREY = '#EEF2EB', '#23412D', '#2C332D', '#6B736C'
CARD, SOFT, ACCENT = '#FFFFFF', '#DCE6D8', '#7A9A6E'
FONT = 'Microsoft JhengHei'


def rgb(h):
    return RGBColor.from_string(h.lstrip('#'))


# ---- analysis.txt 全文 → 備忘稿 ----
analysis = (out / 'analysis.txt').read_text(encoding='utf-8')
notes_by_n = {}
for block in re.split(r'\n(?=Slide \d+\n)', analysis.split('---', 1)[1]):
    m = re.match(r'\s*Slide (\d+)\n', block)
    if m:
        keep = [l for l in block.splitlines()[1:] if re.match(r'^(描述|詮釋|理答|反思)：', l) or l.startswith('（')]
        notes_by_n[int(m.group(1))] = '\n\n'.join(keep)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def new_slide(notes=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def text(s, x, y, w, h, runs, size=14, color=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.15, after=0):
    """runs: str，或 [(文字, {size,color,bold}) …] 的段落清單（list of list）。"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    paras = [[(runs, {})]] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing; p.space_after = Pt(after)
        pPr = p._p.get_or_add_pPr()  # 中文避頭尾：不讓「」？。落到行首
        pPr.set('eaLnBrk', '1'); pPr.set('hangingPunct', '1')
        for t, o in para:
            r = p.add_run(); r.text = t
            f = r.font; f.name = FONT; f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold)
            f.color.rgb = rgb(o.get('color', color))
            rPr = r._r.get_or_add_rPr(); rPr.set('lang', 'zh-TW'); rPr.set('altLang', 'en-US')
            ea = rPr.makeelement(qn('a:ea'), {'typeface': FONT})  # 中文字型要設在 ea，否則會退回新細明體
            rPr.append(ea)
    return tb


def footer(s, label):
    text(s, 0.7, 7.02, 9, 0.3, label, size=10, color=GREY)


# ---- 1. 封面 ----
s = new_slide()
box(s, 9.2, -1.2, 5.6, 5.6, fill=SOFT, shape=MSO_SHAPE.OVAL)
box(s, 10.6, 4.4, 3.4, 3.4, fill='#C9D8C3', shape=MSO_SHAPE.OVAL)
box(s, 0.9, 1.55, 0.12, 2.6, fill=GREEN)
text(s, 1.25, 1.4, 9.5, 2.9, [[(l, {})] for l in C['cover']['title'].split('\n')], size=44, color=GREEN, bold=True,
     spacing=1.1)
text(s, 1.25, 4.45, 9.5, 0.6, C['cover']['subtitle'], size=22, color=BODY)
text(s, 1.25, 5.25, 9.8, 0.9, C['cover']['meta'], size=14, color=GREY)

# ---- 2. 概覽 ----
o = C['overview']
s = new_slide()
text(s, 0.7, 0.5, 12, 0.8, o['title'], size=34, color=GREEN, bold=True)
for i, side in enumerate((o['left'], o['right'])):
    x = 0.7 + i * 6.1
    box(s, x, 1.55, 5.8, 4.6, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    box(s, x, 1.55, 5.8, 0.12, fill=GREEN if i == 0 else ACCENT)
    text(s, x + 0.35, 1.9, 5.2, 0.5, side['head'], size=22, color=GREEN, bold=True)
    text(s, x + 0.35, 2.45, 5.2, 0.45, side['sub'], size=15, color=GREY)
    text(s, x + 0.35, 3.15, 5.2, 2.9, [[('●  ', {'color': ACCENT, 'size': 12}), (t, {})] for t in side['items']],
         size=17, after=12)
box(s, 0.7, 6.35, 11.9, 0.55, fill=SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
text(s, 0.9, 6.38, 11.5, 0.5, o['lens'], size=14, color=GREEN, bold=True, anchor=MSO_ANCHOR.MIDDLE,
     align=PP_ALIGN.CENTER)

# ---- 3. 各段 ----
for sec in C['sections']:
    s = new_slide(notes_by_n.get(sec['n']))
    text(s, 0.7, 0.35, 8, 0.35, f"第 {sec['n']} 段　{sec['time']}", size=12, color=ACCENT, bold=True)
    text(s, 0.7, 0.68, 12, 0.75, sec['title'], size=32, color=GREEN, bold=True)
    text(s, 0.7, 1.38, 12, 0.45, sec['sub'], size=16, color=GREY)

    # 左欄：引語卡＋描述／詮釋
    box(s, 0.7, 2.0, 7.55, 1.35, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    box(s, 0.7, 2.0, 0.1, 1.35, fill=GREEN)
    qsize = 19 if len(sec['quote']) <= 24 else 16.5
    text(s, 1.0, 2.08, 7.1, 0.9, [[('「', {'color': ACCENT, 'bold': True}), (sec['quote'], {}),
                                    ('」', {'color': ACCENT, 'bold': True})]], size=qsize, color=BODY, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.0, 2.98, 7.1, 0.3, f"— {sec['who']}｜{sec['time']}｜逐字稿可信度：{sec['conf']}", size=10.5, color=GREY)
    def labeled_block(lab, body, x, y, w, h, size=17.5):
        box(s, x, y + 0.04, 0.72, 0.34, fill=GREEN if lab == '詮釋' else ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            radius=0.3)
        text(s, x, y + 0.04, 0.72, 0.34, lab, size=12, color='#FFFFFF', bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.85, y, w - 0.85, h, body, size=size, spacing=1.3)

    if sec.get('image') == 'board' and board_img.exists():
        # 有板書照片的段落：照片放左欄，描述排在照片右側
        s.shapes.add_picture(str(board_img), Inches(0.7), Inches(3.55), Inches(3.6))
        text(s, 0.7, 5.65, 3.6, 0.3, '31:45 黑板（360p 放大）', size=10, color=GREY)
        labeled_block('描述', sec['desc'], 4.5, 3.5, 3.75, 2.3, size=15)
        labeled_block('詮釋', sec['interp'], 0.7, 6.0, 7.55, 1.0, size=15)
    else:
        labeled_block('描述', sec['desc'], 0.7, 3.6, 7.55, 1.55)
        labeled_block('詮釋', sec['interp'], 0.7, 5.25, 7.55, 1.55)

    # 右欄：理答策略＋反思
    rx, rw = 8.55, 4.1
    ry = 2.0
    n = len(sec['strat'])
    ph = 0.55 + n * 0.62
    box(s, rx, ry, rw, ph, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    text(s, rx + 0.25, ry + 0.12, rw - 0.4, 0.35, '理答策略', size=13, color=GREEN, bold=True)
    for i, (tag, desc) in enumerate(sec['strat']):
        yy = ry + 0.52 + i * 0.62
        text(s, rx + 0.25, yy, rw - 0.4, 0.6, [[(tag, {'bold': True, 'color': GREEN, 'size': 14}),
                                                 ('　' + desc, {'size': 13.5})]], spacing=1.1)
    fy = ry + ph + 0.2
    fh = max(1.0, 6.85 - fy)
    box(s, rx, fy, rw, fh, fill=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    text(s, rx + 0.25, fy + 0.12, rw - 0.4, 0.32, '反思', size=13, color='#CFE0C6', bold=True)
    text(s, rx + 0.25, fy + 0.45, rw - 0.45, fh - 0.5, sec['reflect'], size=16, color='#FFFFFF', spacing=1.25)
    if sec.get('note'):
        footer(s, '※ ' + sec['note'])

# ---- 4. 理答分布 ----
st = C['strategy_summary']
s = new_slide()
text(s, 0.7, 0.5, 12, 0.8, st['title'], size=34, color=GREEN, bold=True)
cd = CategoryChartData()
cd.categories = [c for c, _ in st['counts']]
cd.add_series('段落數', [v for _, v in st['counts']])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.7), Inches(1.5), Inches(6.6), Inches(5.1), cd)
ch = gf.chart
ch.has_legend = False; ch.has_title = False
plot = ch.plots[0]; plot.gap_width = 60; plot.has_data_labels = True
plot.data_labels.font.size = Pt(14); plot.data_labels.font.color.rgb = rgb(GREEN); plot.data_labels.font.bold = True
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.series[0].format.fill.solid(); plot.series[0].format.fill.fore_color.rgb = rgb(ACCENT)
ca = ch.category_axis; ca.reverse_order = True
ca.tick_labels.font.size = Pt(16); ca.tick_labels.font.name = FONT; ca.tick_labels.font.color.rgb = rgb(BODY)
ca.format.line.color.rgb = rgb('#B9C7B4')
va = ch.value_axis; va.visible = False; va.has_major_gridlines = False; va.maximum_scale = 8
box(s, 7.7, 1.6, 4.95, 4.4, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, 8.0, 1.85, 4.4, 4.0, [[('●  ', {'color': ACCENT, 'size': 12}), (t, {})] for t in st['notes']], size=16,
     after=14, spacing=1.2)
text(s, 0.7, 6.7, 12, 0.5, st['caption'], size=11, color=GREY)

# ---- 5. 結語 ----
cl = C['closing']
s = new_slide()
text(s, 0.7, 0.5, 12, 0.8, cl['title'], size=34, color=GREEN, bold=True)
for i, q in enumerate(cl['items']):
    y = 1.65 + i * 1.55
    box(s, 0.7, y, 11.9, 1.3, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    box(s, 0.95, y + 0.25, 0.8, 0.8, fill=GREEN, shape=MSO_SHAPE.OVAL)
    text(s, 0.95, y + 0.25, 0.8, 0.8, str(i + 1), size=24, color='#FFFFFF', bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.05, y + 0.1, 10.3, 1.1, q, size=21, color=BODY, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.7, 6.55, 12, 0.6, cl['foot'], size=12, color=GREY)

name = sys.argv[2] if len(sys.argv) > 2 else 'slides.pptx'
dst = deck_dir / name
prs.save(dst)
print('saved', dst, len(prs.slides), 'slides')

# ---- HTML 版（同內容，鍵盤左右鍵翻頁）----
def esc(t):
    return html.escape(t)


pages = []
import base64
board_b64 = base64.b64encode(board_img.read_bytes()).decode() if board_img.exists() else ''  # 內嵌，HTML 單檔可攜
pages.append(f"""<section class="cover"><div class="bar"></div><h1>{'<br>'.join(esc(l) for l in C['cover']['title'].split(chr(10)))}</h1>
<p class="sub">{esc(C['cover']['subtitle'])}</p><p class="meta">{esc(C['cover']['meta'])}</p></section>""")
cols = ''.join(f"""<div class="card"><h3>{esc(x['head'])}</h3><p class="muted">{esc(x['sub'])}</p><ul>{''.join(f'<li>{esc(i)}</li>' for i in x['items'])}</ul></div>"""
               for x in (o['left'], o['right']))
pages.append(f"""<section><h2>{esc(o['title'])}</h2><div class="two">{cols}</div><p class="lens">{esc(o['lens'])}</p></section>""")
for sec in C['sections']:
    img = (f'<figure><img src="data:image/png;base64,{board_b64}" alt="31:45 黑板"><figcaption>31:45 黑板（360p 放大）</figcaption></figure>'
           if sec.get('image') == 'board' and board_b64 else '')
    strat = ''.join(f'<li><b>{esc(t)}</b>　{esc(d)}</li>' for t, d in sec['strat'])
    note = f'<p class="note">※ {esc(sec["note"])}</p>' if sec.get('note') else ''
    pages.append(f"""<section><p class="kicker">第 {sec['n']} 段　{esc(sec['time'])}</p><h2>{esc(sec['title'])}</h2><p class="muted">{esc(sec['sub'])}</p>
<div class="grid"><div><blockquote>「{esc(sec['quote'])}」<cite>— {esc(sec['who'])}｜{esc(sec['time'])}｜可信度：{esc(sec['conf'])}</cite></blockquote>
<p><span class="tag">描述</span>{esc(sec['desc'])}</p><p><span class="tag dark">詮釋</span>{esc(sec['interp'])}</p></div>
<div>{img}<div class="card"><h4>理答策略</h4><ul class="strat">{strat}</ul></div><div class="reflect"><h4>反思</h4><p>{esc(sec['reflect'])}</p></div></div></div>{note}</section>""")
mx = max(v for _, v in st['counts'])
bars = ''.join(f'<div class="row"><span>{esc(c)}</span><div class="track"><div class="fill" style="width:{v / mx * 100:.0f}%"></div></div><b>{v}</b></div>' for c, v in st['counts'])
pages.append(f"""<section><h2>{esc(st['title'])}</h2><div class="grid"><div class="bars">{bars}</div><div class="card"><ul>{''.join(f'<li>{esc(t)}</li>' for t in st['notes'])}</ul></div></div><p class="note">{esc(st['caption'])}</p></section>""")
pages.append(f"""<section><h2>{esc(cl['title'])}</h2><ol class="qs">{''.join(f'<li>{esc(q)}</li>' for q in cl['items'])}</ol><p class="note">{esc(cl['foot'])}</p></section>""")

css = """:root{--bg:#EEF2EB;--green:#23412D;--body:#2C332D;--grey:#6B736C;--card:#fff;--soft:#DCE6D8;--accent:#7A9A6E}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#18201A;--green:#B9D3AE;--body:#E3E8E0;--grey:#9AA597;--card:#222C24;--soft:#2D3A2F;--accent:#8FB083}}
:root[data-theme="dark"]{--bg:#18201A;--green:#B9D3AE;--body:#E3E8E0;--grey:#9AA597;--card:#222C24;--soft:#2D3A2F;--accent:#8FB083}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--body);font-family:"Microsoft JhengHei","Noto Sans TC",sans-serif}
section{display:none;max-width:1200px;margin:0 auto;padding:40px 16px 90px;min-height:100vh}section.on{display:block}
h1{font-size:clamp(30px,5vw,52px);color:var(--green);line-height:1.2;margin:0 0 20px}h2{font-size:clamp(24px,3.4vw,38px);color:var(--green);margin:4px 0}
h3,h4{color:var(--green);margin:0 0 8px}.muted,.meta,cite,.note,figcaption{color:var(--grey)}.kicker{color:var(--accent);font-weight:700;margin:0}
.cover{padding-top:18vh}.cover .bar{width:8px;height:120px;background:var(--green);float:left;margin-right:24px}.sub{font-size:22px}
.two,.grid{display:grid;grid-template-columns:1fr 1fr;gap:24px;margin-top:24px}.grid{grid-template-columns:1.8fr 1fr}
@media (max-width:760px){.two,.grid{grid-template-columns:1fr}}
.card{background:var(--card);border-radius:14px;padding:18px 20px;margin-bottom:16px}blockquote{background:var(--card);border-left:6px solid var(--green);border-radius:10px;margin:0 0 18px;padding:18px 20px;font-size:20px;font-weight:700}
cite{display:block;font-size:13px;font-weight:400;margin-top:8px;font-style:normal}p{line-height:1.7;font-size:17px}
.tag{display:inline-block;background:var(--accent);color:#fff;border-radius:99px;padding:1px 12px;margin-right:10px;font-size:14px;font-weight:700}.tag.dark{background:var(--green)}
.strat{padding-left:18px}.strat li{margin:6px 0;line-height:1.5}.strat b{color:var(--green)}
.reflect{background:var(--green);color:#fff;border-radius:14px;padding:18px 20px}.reflect h4{color:#CFE0C6}.reflect p{margin:0}
:root[data-theme="dark"] .reflect,:root:not([data-theme="light"]) .reflect{color:#fff}
figure{margin:0 0 12px}img{width:100%;border-radius:10px}.lens{background:var(--soft);border-radius:99px;padding:10px 20px;text-align:center;color:var(--green);font-weight:700}
.bars .row{display:grid;grid-template-columns:60px 1fr 30px;align-items:center;gap:10px;margin:14px 0;font-size:18px}.track{background:var(--soft);border-radius:6px;height:26px}.fill{background:var(--accent);height:100%;border-radius:6px}
.qs li{background:var(--card);border-radius:14px;padding:22px;margin:16px 0;font-size:21px;line-height:1.5}
nav{position:fixed;bottom:0;left:0;right:0;display:flex;justify-content:center;gap:16px;align-items:center;padding:12px;background:var(--bg)}
nav button{background:var(--green);color:var(--bg);border:0;border-radius:99px;padding:8px 18px;font-size:15px;cursor:pointer}nav button:disabled{opacity:.35}"""
js = """const S=[...document.querySelectorAll('section')];let i=0;const c=document.getElementById('c'),p=document.getElementById('p'),n=document.getElementById('n');
function go(k){i=Math.max(0,Math.min(S.length-1,k));S.forEach((s,j)=>s.classList.toggle('on',j===i));c.textContent=(i+1)+' / '+S.length;p.disabled=i===0;n.disabled=i===S.length-1;window.scrollTo(0,0)}
p.onclick=()=>go(i-1);n.onclick=()=>go(i+1);document.addEventListener('keydown',e=>{if(e.key==='ArrowRight'||e.key===' ')go(i+1);if(e.key==='ArrowLeft')go(i-1)});go(0);"""
page = f"""<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>立體形體課例研究</title><style>{css}</style></head><body>{''.join(pages)}
<nav><button id="p">上一頁</button><span id="c"></span><button id="n">下一頁</button></nav><script>{js}</script></body></html>"""
(dst.with_suffix('.html')).write_text(page, encoding='utf-8')
print('saved', dst.with_suffix('.html'))
