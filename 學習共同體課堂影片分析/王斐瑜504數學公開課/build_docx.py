"""把 output/analysis.txt 與校正後字幕合成 Word 課例分析報告。

每一節依其時間碼，從 subtitles_final.srt 擷取「已校正」的字幕片段附在描述—詮釋—反思之後；
未校正的辨識結果不收錄（只註明），避免把 Whisper 的錯字當成課堂證據。

用法：python -X utf8 -I build_docx.py <lesson_dir> [輸出檔名]
"""
import json, re, sys
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

root = Path(sys.argv[1])
out = root / 'output'
FONT = 'Microsoft JhengHei'
GREEN, BODY, GREY, BG = RGBColor(0x23, 0x41, 0x2D), RGBColor(0x2C, 0x33, 0x2D), RGBColor(0x6B, 0x73, 0x6C), 'EEF2EB'


def ts2s(t):
    h, m, s = t.replace(',', '.').split(':')
    return int(h) * 3600 + int(m) * 60 + float(s)


def mmss(x):
    x = int(x)
    return f'{x // 60:02}:{x % 60:02}'


# ---- 讀分析稿 ----
text = (out / 'analysis.txt').read_text(encoding='utf-8')
header = [l[2:].strip() for l in text.split('---')[0].splitlines() if l.startswith('# ')]
slides = []
for block in re.split(r'\n(?=Slide \d+)', text.split('---', 1)[1]):
    block = block.strip()
    if not block.startswith('Slide'):
        continue
    s = {'notes': []}
    key = None
    for line in block.splitlines()[1:]:
        m = re.match(r'^(標題|副標題|時間碼|描述|詮釋|理答|反思)：(.*)$', line)
        if m:
            key = m.group(1); s[key] = m.group(2).strip()
        elif line.startswith('（') and line.strip():
            s['notes'].append(line.strip())
        elif line.strip() and key:
            s[key] += line.strip()
    a, b = re.findall(r'\d\d:\d\d:\d\d', s['時間碼'])
    s['range'] = (ts2s(a), ts2s(b))
    slides.append(s)

# ---- 讀字幕與校正範圍 ----
cues = []
for block in re.split(r'\n\s*\n', (out / 'subtitles_final.srt').read_text(encoding='utf-8').strip()):
    lines = block.strip().splitlines()
    a, b = [x.strip() for x in lines[1].split('-->')]
    cues.append((ts2s(a), ts2s(b), ' '.join(lines[2:])))
corr = json.loads((out / 'corrections.json').read_text(encoding='utf-8'))
seg_conf = [(ts2s(a), ts2s(b), c) for a, b, c, _ in corr['segments']]
unclear = [(ts2s(a), ts2s(b)) for a, b, _ in corr['unclear_ranges']]


def conf_of(t):
    for i, (a, b, c) in enumerate(seg_conf):
        if a <= t < b:
            return i, c
    return None


def excerpt(a, b):
    rows, skipped, hit_unclear = [], 0, False
    for s, e, t in cues:
        if not (a <= s < b):
            continue
        hit = conf_of(s)
        if hit:
            seg, c = hit
            if rows and rows[-1][3] == seg:
                rows[-1][2] += t  # 同一校正段的分行字幕合回一句
            else:
                rows.append([s, c, t, seg])
        elif t.startswith('（'):
            hit_unclear = True
        else:
            skipped += 1
    return rows, skipped, hit_unclear


# ---- 版面工具 ----
doc = Document()
for sec in doc.sections:
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)

st = doc.styles['Normal']
st.font.name = FONT; st.font.size = Pt(11); st.font.color.rgb = BODY
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
st.paragraph_format.line_spacing = 1.35
st.paragraph_format.space_after = Pt(4)
for name, size in (('Title', 22), ('Heading 1', 15), ('Heading 2', 12)):
    h = doc.styles[name]
    h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = GREEN
    h.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def run(p, t, bold=False, color=None, size=None, italic=False):
    r = p.add_run(t)
    r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    return r


def labeled(label, body):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0)
    run(p, f'【{label}】', bold=True, color=GREEN)
    run(p, body)


def answering(body):
    """詮釋層的理答策略分析：依【策略】標記拆成條列，策略名加粗。"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    run(p, '理答策略分析', bold=True, color=GREEN, size=10.5)
    pieces = re.split(r'(【[^】]+】)', body)
    lead = pieces[0].strip()
    if lead:
        q = doc.add_paragraph(); q.paragraph_format.left_indent = Cm(0.6)
        run(q, lead, size=10.5)
    for tag, txt in zip(pieces[1::2], pieces[2::2]):
        q = doc.add_paragraph(style='List Bullet')
        q.paragraph_format.left_indent = Cm(1.2)
        run(q, tag.strip('【】') + '　', bold=True, color=GREEN, size=10.5)
        run(q, txt.strip(), size=10.5)


# ---- 封面資訊 ----
doc.add_paragraph(header[0].replace('課例研究分析報告', ''), style='Title')
p = doc.add_paragraph(); run(p, '課例研究分析報告（含校正逐字稿片段）', color=GREEN, size=13)
info = [l for l in header[1:] if not l.startswith('引用說明') and not l.startswith('學生以')]
tbl = doc.add_table(rows=0, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.style = 'Table Grid'
for l in info:
    k, v = l.split('：', 1)
    cells = tbl.add_row().cells
    cells[0].width, cells[1].width = Cm(3.2), Cm(13.4)
    shade(cells[0], BG)
    run(cells[0].paragraphs[0], k, bold=True, color=GREEN)
    cells[1].paragraphs[0].add_run(v)

doc.add_heading('閱讀說明', level=2)
for t in ['本報告依「描述－詮釋－反思」三階層分析，聚焦學生的學習事實，不對授課教師評分或給建議。',
          '每一節後附「逐字稿片段」：取自校正後字幕 subtitles_final.srt 中該節時間範圍內、經三源交叉比對校正的語句，並標可信度（高＝三源一致；中＝兩源一致或依教材語境推定）。',
          '（ ）內為依上下文補足、（…）為聽不清楚；未校正的辨識結果不收錄。學生一律以角色或代稱稱呼。',
          '「理答策略分析」位於各節詮釋層之下，以六大理答策略（轉引、轉問、反問、提示、釐清、深究）分析教師語句，另記等待、代答、評價語與回到文本；只分析校正逐字稿中有的教師語句，描述其作用，不評好壞。',
          '影片為單機位 360p，學生小組討論與各組發表的原音大多無法辨識，相關段落依教師轉述與板書推定，文中已註明。']:
    doc.add_paragraph(t, style='List Bullet')

# ---- 各節 ----
for i, s in enumerate(slides, 1):
    doc.add_page_break() if i == 1 else None
    doc.add_heading(f'{i}．{s["標題"]}', level=1)
    p = doc.add_paragraph()
    run(p, s['副標題'], italic=True, color=GREY)
    run(p, f'　｜　時間碼 {s["時間碼"]}', color=GREY, size=10)
    labeled('描述', s['描述'])
    labeled('詮釋', s['詮釋'])
    if s.get('理答'):
        answering(s['理答'])
    labeled('反思', s['反思'])
    for n in s['notes']:
        p = doc.add_paragraph(); run(p, n, italic=True, color=GREY, size=10)

    a, b = s['range']
    if b - a > 1800:  # 總結節涵蓋整堂課，不重複列逐字稿
        continue
    rows, skipped, hit_unclear = excerpt(a, b)
    doc.add_heading('逐字稿片段', level=2)
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.rows[0].cells[0]; cell.width = Cm(16.6); shade(cell, BG)
    first = True
    for start, c, t, _ in rows:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(3)
        run(p, f'[{mmss(start)}] ', bold=True, color=GREEN, size=10)
        run(p, t, size=10.5)
        run(p, f'　（可信度：{c}）', color=GREY, size=9)
    notes = []
    if hit_unclear:
        notes.append('本時段含學生操作或發表，背景嘈雜，原音無法辨識')
    if skipped:
        notes.append(f'另有 {skipped} 條未經校正的辨識語句未收錄')
    if notes or first:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        run(p, '※ ' + '；'.join(notes or ['本時段無已校正語句']) + '。', italic=True, color=GREY, size=9)

# ---- 頁尾頁碼 ----
for sec in doc.sections:
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        r = p.add_run()
        if tag:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = txt
        r._r.append(el)

dst = out / (sys.argv[2] if len(sys.argv) > 2 else '課例分析_王斐瑜504數學_立體形體.docx')
doc.save(dst)
print('saved', dst, len(slides), 'sections')
